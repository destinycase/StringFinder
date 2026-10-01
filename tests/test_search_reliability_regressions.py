"""Reproductions of silent omissions found in the 5.9.24 reliability review."""
import json
from unittest.mock import patch
import zipfile

import pytest
from core import search_engine as se
from core.json_policy import DuplicateKeyError, ObjectPairs, loads_document
from sf_utils.app_strings import AppStrings
from sf_utils.constants import Constants


@pytest.fixture(autouse=True)
def search_policy():
    with se.use_search_settings_snapshot({Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS: False}):
        yield


@pytest.mark.parametrize("content,query", [(' {"v":"NEEDLE"}', "needle"), ('{"v":"\\u006e\\u0065\\u0065\\u0064\\u006c\\u0065"}', "needle"), ('{"v":"ß"}', "ss"), ('null', "null")])
@pytest.mark.parametrize("existence", [False, True])
def test_precise_json_existence_matches_decoded_values(tmp_path, content, query, existence):
    path = tmp_path / "sample.json"
    path.write_text(content, encoding="utf-8")
    result = se.search_in_json_special(str(path), query, use_complex_search=True, existence_only=existence)
    assert result and result[0] == str(path) and result[1] == 1


@pytest.mark.parametrize("precise", [False, True])
@pytest.mark.parametrize("existence", [False, True])
@pytest.mark.parametrize("allow", [False, True])
def test_duplicate_json_keys_policy(tmp_path, precise, existence, allow):
    path = tmp_path / "duplicate.json"
    path.write_text('{"v":"needle","\\u0076":"needle"}', encoding="utf-8")
    with se.use_search_settings_snapshot({Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS: allow}):
        result = se.search_in_json_special(str(path), "needle", use_complex_search=precise, existence_only=existence)
    assert result
    if allow:
        assert result[0] == str(path) and result[1] == (1 if existence else 2)
    else:
        assert result[0] == Constants.STATUS_SKIPPED
        assert AppStrings.JSON_DETAIL_DUPLICATE_KEYS in result[1]


def test_precise_deep_json_keeps_unicode_and_duplicates():
    content = '[' * 3000 + '{"v":"ß","v":"needle"}' + ']' * 3000
    with pytest.raises(DuplicateKeyError):
        loads_document(content)
    data = loads_document(content, True)
    for _ in range(3000):
        data = data[0]
    assert isinstance(data, ObjectPairs) and data == [("v", "ß"), ("v", "needle")]


def test_precise_deep_json_search_keeps_unicode(tmp_path):
    path = tmp_path / "deep.json"
    path.write_text('[' * 3000 + '"ß"' + ']' * 3000, encoding="utf-8")
    assert se.search_in_json_special(str(path), "ss", use_complex_search=True)[1] == 1


def test_utf16_is_not_excluded_as_binary(tmp_path):
    path = tmp_path / "utf16.txt"
    path.write_bytes("needle 한글".encode("utf-16"))
    assert not se.is_binary_file(str(path))
    assert se.search_in_file(str(path), "needle", use_complex_search=True, exclude_binary=True)[1] == 1


@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be"])
@pytest.mark.parametrize("query", ["le", "needle"])
def test_short_no_bom_utf16_is_text_in_all_paths(tmp_path, encoding, query):
    path = tmp_path / "short.txt"
    path.write_bytes(query.encode(encoding))
    for precise in (False, True):
        assert se.search_in_file(str(path), query, use_complex_search=precise, exclude_binary=True)
    assert se.find_files_with_keyword_fast([str(tmp_path)], query, ["txt"], exclude_binary=True)


def test_scanner_reports_inaccessible_and_missing_roots(tmp_path):
    scanner = se.FileScanner([str(tmp_path)], [])
    with patch.object(se.os, "scandir", side_effect=PermissionError("denied")):
        assert scanner.scan() == []
    assert scanner.skipped and scanner.skipped[0][0] == str(tmp_path)
    missing = se.FileScanner([str(tmp_path / "missing")], [])
    assert missing.scan() == [] and missing.skipped


def test_long_line_detail_includes_match_after_whitespace_and_casefold():
    for prefix in [" " * 5000, "ß" * 5000, "가" * 5000]:
        rows, count, _, _ = se._search_text_stream([prefix + "needle" + "x" * 5000], "needle", exact_match=False, existence_only=False, stop_event=None, max_per_file=10)
        assert count == 1 and "needle" in rows[0][1]


@pytest.mark.parametrize("fragment", ["nee<![CDATA[dle]]>", "nee<!--comment-->dle", "nee<?test abc?>dle"])
@pytest.mark.parametrize("existence", [False, True])
def test_native_xml_matches_adjacent_character_data(tmp_path, fragment, existence):
    path = tmp_path / "split.xml"
    path.write_text(f"<root>{fragment}</root>", encoding="utf-8")
    bits = Constants.RUST_MODE_XML | (Constants.RUST_MODE_EXISTENCE_ONLY if existence else 0)
    results = se.sf_engine.search_file(str(path), "needle", bits)
    assert results
    if not existence:
        assert results[0].offset is None


def test_native_utf8_sampling_boundary(tmp_path):
    path = tmp_path / "boundary.txt"
    path.write_text("a" * 65534 + "가 needle", encoding="utf-8")
    assert se.sf_engine.search_file(str(path), "가", 0)


def test_native_unicode_filename_filter(tmp_path):
    from core.search_engine import search_directory_fast
    path = tmp_path / "Ä-report.txt"
    path.write_text("needle")
    result = search_directory_fast([str(tmp_path)], "needle", ["txt"], filename_filter=["ä"], exclude_hidden=False)
    assert result[Constants.PAYLOAD_RESULTS]


def test_excel_coordinates_and_sheet_failure(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "cells.xlsx"
    book = Workbook()
    book.active["D5"] = "needle"
    book.create_sheet("broken")["A1"] = "needle"
    book.save(path)
    result = se.search_in_excel_special(str(path), "needle", use_complex_search=True)
    assert result and result[1] == 2
    assert ("Sheet", "D5") in [(row[1], row[2]) for row in result[2]]
    corrupt = tmp_path / "corrupt.xlsx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(corrupt, "w") as destination:
        for name in source.namelist():
            if name != "xl/worksheets/sheet1.xml":
                destination.writestr(name, source.read(name))
    result = se.search_in_excel_special(str(corrupt), "needle", use_complex_search=True)
    assert result and (result[0] == Constants.STATUS_SKIPPED or any(row[0] == -2 for row in result[2]))


@pytest.mark.parametrize("precise", [False, True])
def test_json_exact_trims_only_outer_whitespace(tmp_path, precise):
    path = tmp_path / "exact.json"
    path.write_text(json.dumps({"value": " needle "}))
    assert se.search_in_json_special(str(path), "needle", exact_match=True, use_complex_search=precise)


@pytest.mark.parametrize("precise", [False, True])
def test_excel_exact_keeps_internal_whitespace(tmp_path, precise):
    from openpyxl import Workbook
    path = tmp_path / "spaces.xlsx"
    book = Workbook()
    book.active["A1"] = "a b"
    book.save(path)
    assert se.search_in_excel_special(str(path), "ab", exact_match=True, use_complex_search=precise) is None
    assert se.search_in_excel_special(str(path), "a b", exact_match=True, use_complex_search=precise)


def test_native_git_exclude_is_not_a_search_filter(tmp_path):
    import subprocess
    subprocess.run(["git", "init", "-q", str(tmp_path)], check=True)
    (tmp_path / ".git" / "info" / "exclude").write_text("excluded.txt\n")
    (tmp_path / "excluded.txt").write_text("needle")
    result = se.search_directory_fast([str(tmp_path)], "needle", ["txt"], exclude_hidden=False)
    assert len(result[Constants.PAYLOAD_RESULTS]) == 1


@pytest.mark.parametrize("existence", [False, True])
@pytest.mark.parametrize("allow", [False, True])
def test_duplicate_policy_in_native_directory(tmp_path, existence, allow):
    (tmp_path / "data.json").write_text('{"v":"needle","v":"other"}')
    with se.use_search_settings_snapshot({Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS: allow}):
        result = se.search_directory_fast([str(tmp_path)], "needle", ["json"], special_mode=Constants.MODE_JSON, existence_only=existence)
    assert bool(result[Constants.PAYLOAD_RESULTS]) is allow
    assert bool(result[Constants.PAYLOAD_SKIPPED]) is not allow


def test_duplicate_check_validates_after_limits_and_hits(tmp_path):
    path = tmp_path / "after-hit.json"
    path.write_text('{"v":"needle","child":{"x":1,"x":2}}')
    for existence in (False, True):
        with pytest.raises(RuntimeError, match="duplicate JSON object key"):
            se.sf_engine.search_file(str(path), "needle", Constants.RUST_MODE_JSON | (Constants.RUST_MODE_EXISTENCE_ONLY if existence else 0), max_per_file=1, max_json_depth=1)


def test_model_does_not_silently_drop_files_at_100000(qtbot):
    from ui.models import SearchResultModel
    model = SearchResultModel()
    model.add_results([(f"/file-{index}.txt", 1, []) for index in range(100001)])
    assert len(model._result_buffer) == 100001


def test_failed_search_final_summary_is_not_complete(qtbot, mock_config_manager):
    import time
    from ui.search_tab import SearchTab
    tab = SearchTab(mock_config_manager)
    qtbot.addWidget(tab)
    tab.scan_start_time = time.time()
    tab.search_state = Constants.SearchState.SEARCHING
    tab._on_search_error("Injected engine failure")
    tab._on_search_finished(0, 0, 0)
    assert AppStrings.SUMMARY_PREFIX_FAILED in tab.result_view_panel.summary_label.text()
    assert AppStrings.SUMMARY_PREFIX_FINISHED not in tab.result_view_panel.summary_label.text()


def test_python_worker_reports_scanner_errors(tmp_path):
    from core.worker import SearchWorker
    worker = SearchWorker({Constants.PAYLOAD_SEARCH_PATHS: [str(tmp_path)], Constants.PAYLOAD_SEARCH_STRING: "needle", Constants.PAYLOAD_USE_COMPLEX_SEARCH: True})
    skipped, finished = [], []
    worker.signals.skipped_found.connect(skipped.extend)
    worker.signals.search_finished.connect(lambda *counts: finished.append(counts))
    with patch.object(se.os, "scandir", side_effect=PermissionError("denied")):
        worker._run_python_search(force_python=True)
    assert skipped and finished == [(0, 0, 1)]


@pytest.mark.parametrize("content", ['{"a":1,}', '[1,]', '{"a" 1}', '[true false]', 'null null', '[', '{"a":NaN}'])
def test_stack_safe_decoder_rejects_invalid_grammar(content):
    with patch("core.json_policy.json.loads", side_effect=RecursionError):
        with pytest.raises(ValueError):
            loads_document(content)


@pytest.mark.parametrize("value", [None, 1, "true", [], {}])
def test_duplicate_setting_invalid_type_uses_default(mock_config_manager, value):
    normalized = mock_config_manager._normalize_advanced_settings({Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS: value})
    assert normalized[Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS] is False


@pytest.mark.parametrize("allow", [False, True])
@pytest.mark.parametrize("entrypoint", ["search_files_list", "find_filenames"])
def test_duplicate_policy_other_native_entrypoints(tmp_path, allow, entrypoint):
    path = tmp_path / "duplicate.json"
    path.write_text('{"v":"needle","v":"other"}')
    bits = Constants.RUST_MODE_JSON | (Constants.RUST_MODE_ALLOW_DUPLICATE_JSON_KEYS if allow else 0)
    if entrypoint == "search_files_list":
        found, skipped = se.sf_engine.search_files_list(file_list=[str(path)], search_string="needle", mode_bits=bits)
    else:
        found, skipped = se.sf_engine.find_files_with_keyword(paths=[str(tmp_path)], keyword="needle", extensions=["json"], mode_bits=bits)
    assert bool(found) is allow and bool(skipped) is not allow


def test_native_duplicate_detection_after_inline_keys_overflow(tmp_path):
    path = tmp_path / "many-keys.json"
    path.write_text('{' + ','.join(f'"key{i}":"needle"' for i in range(20)) + ',"key0":"other"}')
    with pytest.raises(RuntimeError, match="duplicate JSON object key"):
        se.sf_engine.search_file(str(path), "needle", Constants.RUST_MODE_JSON)


def test_duplicate_setting_combo_persists_and_resets(qtbot, mock_config_manager):
    from ui.settings_dialog import SettingsDialog
    dialog = SettingsDialog(mock_config_manager)
    qtbot.addWidget(dialog)
    combo = dialog.allow_duplicate_json_keys_combo
    assert combo.currentData() is False
    combo.setCurrentIndex(1)
    assert mock_config_manager.get_advanced_settings()[Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS] is True
    with patch("ui.settings_dialog.QMessageBox.information"):
        dialog._reset_advanced_settings()
    assert combo.currentData() is False
    assert mock_config_manager.get_advanced_settings()[Constants.CONFIG_KEY_ALLOW_DUPLICATE_JSON_KEYS] is False


@pytest.mark.parametrize("allow", [False, True])
def test_stack_safe_decoder_matches_standard_decoder(allow):
    """Exercise both decoder implementations with a reproducible varied corpus."""
    import random
    rng = random.Random(5924)

    def value(depth):
        scalar = rng.choice([None, True, False, -42, 1.25, "", "ß needle", "a\n\"b", "가"])
        if depth == 0:
            return scalar
        choice = rng.randrange(3)
        if choice == 0:
            return [value(depth - 1) for _ in range(rng.randrange(5))]
        if choice == 1:
            return {f"key{index}": value(depth - 1) for index in range(rng.randrange(5))}
        return scalar

    for _ in range(200):
        content = json.dumps(value(4), ensure_ascii=rng.choice([True, False]))
        expected = loads_document(content, allow)
        with patch("core.json_policy.json.loads", side_effect=RecursionError):
            assert loads_document(content, allow) == expected


@pytest.mark.parametrize("precise", [False, True])
def test_duplicate_keys_in_separate_objects_are_not_duplicates(tmp_path, precise):
    path = tmp_path / "separate.json"
    path.write_text('[{"v":"needle"},{"v":"needle"}]', encoding="utf-8")
    result = se.search_in_json_special(str(path), "needle", use_complex_search=precise)
    assert result and result[0] == str(path) and result[1] == 2


@pytest.mark.parametrize("precise", [False, True])
def test_duplicate_policy_does_not_restrict_plain_text_search(tmp_path, precise):
    path = tmp_path / "plain.json"
    path.write_text('{"v":"needle","v":"needle"}', encoding="utf-8")
    result = se.search_in_file(str(path), "needle", use_complex_search=precise)
    assert result and result[0] == str(path)


def test_unexpected_worker_failure_finalizes_partial_results(qtbot, mock_config_manager):
    import time
    from core.worker import SearchWorker
    from ui.search_tab import SearchTab
    tab = SearchTab(mock_config_manager)
    qtbot.addWidget(tab)
    tab.scan_start_time = time.time()
    tab.search_state = Constants.SearchState.SEARCHING
    worker = SearchWorker({Constants.PAYLOAD_SEARCH_PATHS: ["unused"], Constants.PAYLOAD_SEARCH_STRING: "needle", Constants.PAYLOAD_USE_COMPLEX_SEARCH: True})
    tab.worker = worker
    worker.signals.error.connect(tab._on_search_error)
    worker.signals.finished.connect(tab._on_worker_finished)
    worker.signals.search_finished.connect(tab._on_search_finished)
    worker.signals.results_found.connect(tab._on_results_found)

    def fail():
        worker.signals.results_found.emit([("/partial.txt", 1, [(1, "needle", None, None)])])
        raise RuntimeError("Injected unexpected worker failure")

    with patch.object(worker, "_run_python_search", side_effect=fail):
        worker.run()
    assert tab.search_state == Constants.SearchState.IDLE
    assert AppStrings.SUMMARY_PREFIX_FAILED in tab.result_view_panel.summary_label.text()
    assert tab.result_view_panel.result_model.rowCount() == 1


def test_benchmark_waits_for_runnable_cleanup(qtbot, tmp_path):
    from PySide6.QtCore import QThreadPool
    from scripts import benchmark_performance as benchmark
    calls = []

    class TrackingPool(QThreadPool):
        def waitForDone(self, *args):
            calls.append(True)
            return super().waitForDone(*args)

    path = tmp_path / "small.txt"
    path.write_text("needle", encoding="utf-8")
    with patch.object(benchmark, "QThreadPool", TrackingPool):
        result = benchmark.run_benchmark("cleanup", path, "needle")
    assert calls == [True]
    assert result["results_count"] == 1 and result["skipped_count"] == 0


@pytest.mark.parametrize("allow", [False, True])
def test_native_decoded_duplicate_after_hash_transition(tmp_path, allow):
    path = tmp_path / "escaped-wide.json"
    path.write_text('{' + ','.join(f'"key{i}":"needle"' for i in range(12))
                    + ',"\\u006bey9":"needle"}', encoding="utf-8")
    bits = Constants.RUST_MODE_JSON | (Constants.RUST_MODE_ALLOW_DUPLICATE_JSON_KEYS if allow else 0)
    if allow:
        assert len(se.sf_engine.search_file(str(path), "needle", bits)) == 13
    else:
        with pytest.raises(RuntimeError, match="duplicate JSON object key"):
            se.sf_engine.search_file(str(path), "needle", bits)


def test_native_borrowed_keys_in_deep_object_paths(tmp_path):
    path = tmp_path / "deep-object.json"
    path.write_text('{"key":' * 300 + '"needle"' + '}' * 300, encoding="utf-8")
    result = se.sf_engine.search_file(str(path), "needle", Constants.RUST_MODE_JSON, max_json_depth=1000)
    assert len(result) == 1
    assert result[0].content == "/key" * 300 + "\tneedle"


def test_json_policy_benchmark_checks_both_policies(tmp_path):
    from tools.benchmark_json_policy import measure
    path = tmp_path / "benchmark.json"
    path.write_text('{"v":"needle"}', encoding="utf-8")
    result = measure(se.sf_engine, ("small-sparse", path, "needle", False), 2)
    assert len(result["deny_samples_s"]) == len(result["allow_samples_s"]) == 2
    assert result["deny_median_s"] > 0 and result["allow_median_s"] > 0
    assert len(result["fingerprint"]) == 64


@pytest.mark.parametrize("backend", ["native", "normal"])
def test_json_policy_benchmark_interleaves_engines(tmp_path, backend):
    from tools.benchmark_json_policy import compare
    path = tmp_path / "comparison.json"
    path.write_text('{"v":"needle"}', encoding="utf-8")
    result = compare({"baseline": se.sf_engine, "current": se.sf_engine},
                     ("small-sparse", path, "needle", False), 2, backend)
    assert [row["engine"] for row in result["engines"]] == ["baseline", "current"]
    assert all(len(row["deny_samples_s"]) == len(row["allow_samples_s"]) == 2 for row in result["engines"])
