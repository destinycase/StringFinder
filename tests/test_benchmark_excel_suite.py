"""Validate Excel performance fixtures and fail closed on missing/partial results."""

from pathlib import Path
import logging
import os
import sys

from openpyxl import load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from tools import benchmark_excel_suite as bench  # noqa: E402


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory):
    return bench.create_fixtures(tmp_path_factory.mktemp("excel-suite"), 8, 4)


@pytest.mark.parametrize("kind", range(4))
@pytest.mark.parametrize("case_index", range(16))
def test_native_and_precise_fixture_contract(fixtures, kind, case_index):
    from core.search_engine import HAS_RUST_ENGINE

    if not HAS_RUST_ENGINE:
        pytest.skip("native engine required")
    fixture = fixtures[kind]
    case = bench.cases_for(fixture)[case_index]
    bench.validate_result(bench._run(fixture.path, case), case)


def test_sparse_extent_and_formula_cache_are_real(fixtures):
    sparse, formula = fixtures[1:3]
    book = load_workbook(sparse.path, data_only=True)
    assert book.active.max_row == 100_000
    assert book.active.max_column == 1000
    assert book.active["ALL100000"].value == bench.LAST
    assert sparse.stored_cells == 2 and sparse.rectangle_cells == 100_000_000
    book.close()
    book = load_workbook(formula.path, data_only=True)
    assert book.active["A1"].value == bench.FIRST
    assert book.active["D8"].value == bench.LAST
    assert book.active["B2"].value == 2
    book.close()
    book = load_workbook(formula.path, data_only=False)
    assert book.active["A1"].data_type == "f"
    book.close()


@pytest.mark.parametrize("result", [None, ("skipped", "error"),
    ("book.xlsx", 1, [(0, "Data", "B2", bench.FIRST)]),
    ("book.xlsx", 1, [(0, "Data", "A1", "wrong value")]),
    ("book.xlsx", 1, [(0, "Data", "A1", bench.FIRST), (-2, "error", "", "")])])
def test_validation_rejects_missing_wrong_or_partial_results(result):
    case = bench.ExcelCase("first", bench.FIRST, False, frozenset({"A1"}))
    with pytest.raises(RuntimeError):
        bench.validate_result(result, case)


def test_measure_excludes_warmup_and_validates_each_repeat(monkeypatch):
    calls = []
    case = bench.ExcelCase("absent", "missing", False, frozenset())
    monkeypatch.setattr(bench, "_run", lambda *args: calls.append(args))
    result = bench.measure(Path("unused.xlsx"), case, 3)
    assert len(calls) == 4
    assert len(result["samples"]) == len(result["rss_samples_mib"]) == 3
    assert result["hits"] == 0


def test_failure_prevents_recording(monkeypatch, fixtures):
    case = bench.cases_for(fixtures[0])[0]
    monkeypatch.setattr(bench, "_run", lambda *args: None)
    with pytest.raises(RuntimeError):
        bench.measure(fixtures[0].path, case, 2)


@pytest.mark.parametrize("rows,columns", [(1, 4), (8, 1), (1_048_577, 4), (8, 16_385)])
def test_invalid_dimensions_rejected(tmp_path, rows, columns):
    with pytest.raises(ValueError):
        bench.create_fixtures(tmp_path, rows, columns)


def test_owned_logs_closed_but_other_logs_preserved(tmp_path):
    owned = tmp_path / "owned"
    owned.mkdir()
    logger = logging.getLogger("excel-bench-cleanup-test")
    local = logging.FileHandler(owned / "local.log")
    external = logging.FileHandler(tmp_path / "external.log")
    logger.addHandler(local)
    logger.addHandler(external)
    try:
        bench.close_owned_logs(owned)
        assert local not in logger.handlers and local.stream is None
        assert external in logger.handlers and external.stream is not None
        (owned / "local.log").unlink()
    finally:
        logger.removeHandler(external)
        external.close()


def test_cli_restores_environment_and_removes_temporary_files(monkeypatch, tmp_path):
    directories = []

    def fake_suite(directory, *args):
        directories.append(directory.parent)
        directory.mkdir()
        return {"results": []}

    monkeypatch.setattr(bench, "run_suite", fake_suite)
    monkeypatch.setenv("APPDATA", str(tmp_path / "user-settings"))
    monkeypatch.setattr(sys, "argv", ["benchmark_excel_suite.py"])
    bench.main()
    assert os.environ["APPDATA"] == str(tmp_path / "user-settings")
    assert not directories[0].exists()


def test_cli_restores_environment_on_failure(monkeypatch):
    monkeypatch.delenv("APPDATA", raising=False)
    monkeypatch.setattr(sys, "argv", ["benchmark_excel_suite.py"])

    def fail(*args):
        raise RuntimeError("injected validation failure")

    monkeypatch.setattr(bench, "run_suite", fail)
    with pytest.raises(RuntimeError):
        bench.main()
    assert "APPDATA" not in os.environ
