import re
import pytest
from PySide6.QtWidgets import QApplication, QMessageBox, QFileIconProvider
from sf_utils.constants import Constants
from ui.panels import SearchOptionsPanel
from core.search_engine import search_in_excel_special
from ui.result_view import ResultView
from sf_utils.config_manager import ConfigManager


class TestQAEdgeFixesVerification:
    """수정 버전에 대한 QA 재검증 테스트"""

    def test_search_options_panel_profile_session_persistence(self, qtbot):
        """검색 패널 프로필(정밀 검색 / 존재 확인) 선택 시 get_state 직렬화 정확성 검증"""
        panel = SearchOptionsPanel()
        qtbot.addWidget(panel)

        # 프로필 2: 정밀 검색 (True, False)
        panel.search_profile_combo.setCurrentIndex(2)
        state = panel.get_state()
        assert state[Constants.PAYLOAD_USE_COMPLEX_SEARCH] is True
        assert state[Constants.PAYLOAD_EXISTENCE_ONLY] is False

        # 프로필 3: 정밀 검색 + 존재 확인 (True, True)
        panel.search_profile_combo.setCurrentIndex(3)
        state_both = panel.get_state()
        assert state_both[Constants.PAYLOAD_USE_COMPLEX_SEARCH] is True
        assert state_both[Constants.PAYLOAD_EXISTENCE_ONLY] is True

    def test_excel_complex_search_uppercase_exact_match(self, tmp_path):
        """엑셀 Python 폴백 검색 시 대문자 검색어의 exact_match 탐지 보장 검증"""
        import openpyxl
        wb_path = tmp_path / "test_uppercase.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DataSheet"
        ws["A1"] = "ADMIN_USER"
        ws["B1"] = "admin_user"
        wb.save(wb_path)

        result = search_in_excel_special(
            str(wb_path),
            search_string="ADMIN_USER",
            exact_match=True,
            use_complex_search=True,
        )

        assert result is not None
        assert result[0] != Constants.STATUS_SKIPPED
        file_path, match_count, matches = result
        assert match_count >= 1

    def test_export_results_feedback_called(self, monkeypatch, qtbot):
        """내보내기 실패 시 critical 팝업, 성공 시 information 팝업 호출 검증"""
        icon_provider = QFileIconProvider()
        view = ResultView(icon_provider=icon_provider, config_manager=ConfigManager())
        qtbot.addWidget(view)
        view.set_results([[1, "test.txt", "C:/", "C:/test.txt", [(1, "match")]]])

        # 1. 실패 시나리오
        monkeypatch.setattr(
            "PySide6.QtWidgets.QFileDialog.getSaveFileName",
            lambda *args, **kwargs: ("C:/fake/path.xlsx", "Excel (*.xlsx)")
        )
        def mock_save_fail(*args, **kwargs):
            raise PermissionError("Locked file")
        monkeypatch.setattr("openpyxl.Workbook.save", mock_save_fail)

        critical_called = []
        monkeypatch.setattr(
            QMessageBox, "critical",
            lambda parent, title, msg: critical_called.append((title, msg))
        )
        view._export_results()
        assert len(critical_called) == 1

        # 2. 성공 시나리오
        monkeypatch.setattr("openpyxl.Workbook.save", lambda *args, **kwargs: None)
        info_called = []
        monkeypatch.setattr(
            QMessageBox, "information",
            lambda parent, title, msg: info_called.append((title, msg))
        )
        view._export_results()
        assert len(info_called) == 1
