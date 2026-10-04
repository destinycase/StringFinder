"""Sparse XLSX reading preserves data and precise-search policy."""
from datetime import datetime, time, timedelta
from types import SimpleNamespace
import os
import random
from pathlib import Path
import subprocess
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest

from core import search_engine as se
from core.excel_sparse import iter_sparse_cells
from sf_utils.constants import Constants


def replace_sheet(path, xml):
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, 'w') as archive:
        for info, data in entries:
            archive.writestr(info, xml if info.filename == 'xl/worksheets/sheet1.xml' else data)


@pytest.fixture
def book(tmp_path):
    path = tmp_path / 'cells.xlsx'
    wb = Workbook()
    wb.active.title = 'Data'
    wb.active['C3'] = 'needle first'
    wb.active['E7'] = 'Straße e\u0301'
    wb.save(path)
    return path


def assert_precise_equals_legacy(monkeypatch, path, query, exact=False, existence=False):
    new = se.search_in_excel_special(str(path), query, exact_match=exact,
                                    use_complex_search=True, existence_only=existence)
    with monkeypatch.context() as patch:
        patch.setattr(se, 'HAS_RUST_ENGINE', False)
        old = se.search_in_excel_special(str(path), query, exact_match=exact,
                                        use_complex_search=True, existence_only=existence)
    assert new == old
    return new


@pytest.mark.parametrize('query', ['needle', 'strasse', 'é', 'missing', ' '])
@pytest.mark.parametrize('exact', [False, True])
@pytest.mark.parametrize('existence', [False, True])
def test_precise_sparse_matches_legacy(book, monkeypatch, query, exact, existence):
    assert_precise_equals_legacy(monkeypatch, book, query, exact, existence)


@pytest.mark.parametrize('value', [False, 0, 1e-5, '#DIV/0!', datetime(1904, 1, 1),
                                  datetime(2026, 10, 1, 12, 1, 2, 123000),
                                  time(12, 1, 2), timedelta(days=2, seconds=5)])
@pytest.mark.parametrize('epoch1904', [False, True])
def test_typed_decoding_preserved(tmp_path, monkeypatch, value, epoch1904):
    path = tmp_path / 'typed.xlsx'
    wb = Workbook()
    if epoch1904:
        wb.epoch = CALENDAR_MAC_1904
    wb.active['B3'] = value
    wb.save(path)
    from python_calamine import CalamineWorkbook
    legacy = CalamineWorkbook.from_path(str(path)).get_sheet_by_index(0)
    query = se._excel_cell_text(list(legacy.iter_rows())[-1][0])
    if query:
        assert_precise_equals_legacy(monkeypatch, path, query, True)
    else:
        assert_precise_equals_legacy(monkeypatch, path, '#DIV/0!')


@pytest.mark.parametrize('value', ['2026-10-04T00:00:00Z',
                                  '2026-10-04T12:34:56+09:00',
                                  '2026-10-04T12:34:56.123456-05:30'])
@pytest.mark.parametrize('exact', [False, True])
@pytest.mark.parametrize('existence', [False, True])
def test_precise_timezone_iso_keeps_source_literal(book, monkeypatch, value, exact, existence):
    replace_sheet(book, f'''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
    <sheetData><row r="1"><c r="A1" t="d"><v>{value}</v></c></row></sheetData></worksheet>'''.encode())
    result = assert_precise_equals_legacy(monkeypatch, book, value, exact, existence)
    assert result and result[1] == 1
    if not existence:
        assert result[2][0][3] == value
    # The normal engine's existing ISO policy is intentionally unchanged.
    normal_query = value.replace('T', ' ')
    native = se.search_in_excel_special(str(book), normal_query, exact_match=True)
    assert native and native[1] == 1 and native[2][0][3] == normal_query


@pytest.mark.parametrize('serial', [-1, -.1, 0, .1, .9999999999,
                                   1.9999999999, 59.9999999999,
                                   60.9999999999, 365.9999999999, 1461.9999999999])
@pytest.mark.parametrize('epoch1904', [False, True])
@pytest.mark.parametrize('style', ['yyyy-mm-dd', 'hh:mm:ss.000'])
def test_precise_serial_time_rounding_preserves_legacy(tmp_path, monkeypatch, serial, epoch1904, style):
    path = tmp_path / 'serial.xlsx'
    wb = Workbook()
    if epoch1904:
        wb.epoch = CALENDAR_MAC_1904
    wb.active['A1'] = serial
    wb.active['A1'].number_format = style
    wb.save(path)
    from python_calamine import CalamineWorkbook
    old = CalamineWorkbook.from_path(str(path)).get_sheet_by_index(0)
    query = se._excel_cell_text(list(old.iter_rows())[0][0])
    if epoch1904 and 0 <= serial < 1 and style == 'yyyy-mm-dd':
        query = '1904-01-01' if query == '00:00:00' else '1904-01-01 ' + query
    for exact in (False, True):
        for existence in (False, True):
            result = assert_precise_equals_legacy(monkeypatch, path, query, exact, existence)
            assert result and result[1] == 1
            if not existence:
                assert result[2][0][3] == query


@pytest.mark.parametrize('value', [
    '2026-10-04T12:34:56.123456789', '2026-10-04T00:00:00.000000001',
    '2026-10-04T12:34', '2026-02-30T12:34:56', 'not-a-dateTvalue',
    '2026-10-04T23:59:60', '2026-10-04t12:34:56',
    '12:34:56.123456789', '2026-10-04', '0001-01-01T00:00:00',
    '10000-01-01T00:00:00',
])
@pytest.mark.parametrize('exact', [False, True])
@pytest.mark.parametrize('existence', [False, True])
def test_precise_iso_parse_and_microsecond_contract(book, monkeypatch, value, exact, existence):
    replace_sheet(book, f'''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
    <sheetData><row r="1"><c r="A1" t="d"><v>{value}</v></c></row></sheetData></worksheet>'''.encode())
    from python_calamine import CalamineWorkbook
    old = CalamineWorkbook.from_path(str(book)).get_sheet_by_index(0)
    query = se._excel_cell_text(list(old.iter_rows())[0][0])
    result = assert_precise_equals_legacy(monkeypatch, book, query, exact, existence)
    assert result and result[1] == 1
    if not existence:
        assert result[2][0][3] == query


@pytest.mark.parametrize('epoch1904', [False, True])
@pytest.mark.parametrize('style', ['hh:mm:ss.000', '[h]:mm:ss.000'])
def test_precise_random_serial_conversion_matches_legacy(tmp_path, epoch1904, style):
    path = tmp_path / 'serial_bulk.xlsx'
    wb = Workbook()
    if epoch1904:
        wb.epoch = CALENDAR_MAC_1904
    rng = random.Random(603)
    serials = [66659.81668219328, -2000.1, -.9999999999, .9999999999,
               59.9999999999, 60.9999999999, 61.0000000001]
    serials += [rng.uniform(-1, 100000) for _ in range(400)]
    for index, serial in enumerate(serials, 1):
        wb.active.cell(index, 1, serial).number_format = style
    wb.save(path)
    from python_calamine import CalamineWorkbook
    old = CalamineWorkbook.from_path(str(path)).get_sheet_by_index(0)
    expected = [se._excel_cell_text(row[0]) for row in old.iter_rows()]
    actual = se.sf_engine.SparseExcelWorkbook(str(path)).read_sheet('Sheet')[0]
    assert actual == [(index, 0, text) for index, text in enumerate(expected)]


def test_precise_1904_millisecond_rounding_search(tmp_path, monkeypatch):
    path = tmp_path / 'millisecond.xlsx'
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    wb.active['A1'] = 66659.81668219328
    wb.active['A1'].number_format = 'hh:mm:ss.000'
    wb.save(path)
    query = '2086-07-03 19:36:01.342000'
    for exact in (False, True):
        for existence in (False, True):
            result = assert_precise_equals_legacy(monkeypatch, path, query, exact, existence)
            assert result and result[1] == 1


def test_duplicate_out_of_order_cells_last_value_wins(book, monkeypatch):
    replace_sheet(book, b'''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>
    <row r="7"><c r="E7" t="inlineStr"><is><t>needle old</t></is></c></row>
    <row r="3"><c r="C3" t="inlineStr"><is><t>needle first</t></is></c></row>
    <row r="7"><c r="E7" t="inlineStr"><is><t>needle last</t></is></c><c r="E7"/></row>
    </sheetData></worksheet>''')
    result = assert_precise_equals_legacy(monkeypatch, book, 'needle')
    assert [m[2:] for m in result[2]] == [('C3', 'needle first'), ('E7', 'needle last')]
    native = se.search_in_excel_special(str(book), 'needle')
    assert [m[2:4] for m in native[2]] == [('C3', 'needle first'), ('E7', 'needle last')]


def test_malformed_sheet_is_not_reported_as_complete(book, monkeypatch):
    replace_sheet(book, b'<worksheet><sheetData><row><c r="A1" t="inlineStr"><is><t>needle</t></is></c>')
    for precise in (False, True):
        result = se.search_in_excel_special(str(book), 'needle', use_complex_search=precise)
        assert result and result[0] == Constants.STATUS_SKIPPED


def test_empty_query_adapter_fills_gaps_lazily():
    wb = SimpleNamespace(read_sheet=lambda _: ([(2, 4, 'value')], (2, 4), (2, 4)))
    assert list(iter_sparse_cells(wb, 'Data', True)) == [(0, 4, ''), (1, 4, ''), (2, 4, 'value')]


def test_sparse_decoder_used_in_precise_mode(book, monkeypatch):
    assert isinstance(getattr(se.sf_engine, 'SparseExcelWorkbook', None), type)
    def forbidden(*args):
        raise AssertionError('Rectangular parser must not be used for XLSX')
    monkeypatch.setattr('python_calamine.CalamineWorkbook.from_path', forbidden)
    result = se.search_in_excel_special(str(book), 'needle', use_complex_search=True)
    assert result and result[1] == 1


def test_empty_query_stops_enumerating_after_display_capacity(book, monkeypatch):
    monkeypatch.setattr(se, '_positive_limit', lambda *args: 1)
    result = assert_precise_equals_legacy(monkeypatch, book, ' ')
    assert result and result[1] == 2 and result[2][1][0] == -1


@pytest.mark.skipif(sys.platform != 'win32', reason='Windows Job Object isolation')
def test_huge_blank_rectangle_searches_under_256_mib(book, tmp_path):
    from windows_memory_job import memory_job
    replace_sheet(book, b'''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>
    <row r="1"><c r="A1" t="inlineStr"><is><t>needle first</t></is></c></row>
    <row r="100000"><c r="ALL100000" t="inlineStr"><is><t>needle last</t></is></c></row>
    </sheetData></worksheet>''')
    code = '''
import sys
from windows_memory_job import attach_current_process
attach_current_process(sys.argv[1])
from core import search_engine as se
from sf_utils.constants import Constants as C
with se.use_search_settings_snapshot({C.CONFIG_KEY_MAX_PER_FILE_MATCHES: 10000}):
    for precise in (False, True):
        for existence in (False, True):
            for text, coordinate in [('needle first','A1'), ('needle last','ALL100000')]:
                result = se.search_in_excel_special(sys.argv[2], text, use_complex_search=precise, existence_only=existence)
                assert result and result[1] == 1, result
                if not existence:
                    assert result[2][0][2] == coordinate, result
            assert se.search_in_excel_special(sys.argv[2], 'missing', use_complex_search=precise, existence_only=existence) is None
print('BOUNDED_SEARCH_PASS')
'''
    root = Path(__file__).resolve().parents[1]
    environment = dict(os.environ, PYTHONPATH=os.pathsep.join([str(root/'tests'), str(root/'src')]),
                       APPDATA=str(tmp_path/'settings'))
    with memory_job(256) as job:
        result = subprocess.run([sys.executable, '-c', code, job, str(book)], capture_output=True,
                                text=True, timeout=15, env=environment)
    assert result.returncode == 0, result.stderr
    assert 'BOUNDED_SEARCH_PASS' in result.stdout
