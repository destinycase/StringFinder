"""Complement A-J with XLSX shape/policy benchmarks, excluding GUI/worker startup."""

from __future__ import annotations

import argparse
import gc
import hashlib
import json
import logging
import os
import platform
import statistics
import sys
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import psutil
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
FIRST, LAST = "sf_excel_hit_first", "sf_excel_hit_last"


@dataclass(frozen=True)
class ExcelFixture:
    name: str
    path: Path
    first_query: str
    last_query: str
    both_query: str
    first_cell: str
    last_cell: str
    stored_cells: int
    rectangle_cells: int


@dataclass(frozen=True)
class ExcelCase:
    name: str
    query: str
    existence_only: bool
    expected_cells: frozenset[str]
    precise: bool = False


def _cache_formulas(path: Path, first_cell: str, last_cell: str) -> None:
    """Inject known cached results: openpyxl cannot calculate formulas."""
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for info, data in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{NS}}}c"):
                    if cell.find(f"{{{NS}}}f") is None:
                        continue
                    value = cell.find(f"{{{NS}}}v")
                    if value is None:
                        value = ET.SubElement(cell, f"{{{NS}}}v")
                    coordinate = cell.attrib["r"]
                    cell.set("t", "str" if coordinate in (first_cell, last_cell) else "n")
                    value.text = FIRST if coordinate == first_cell else LAST if coordinate == last_cell else "2"
                data = ET.tostring(root, encoding="utf-8")
            archive.writestr(info, data)


def create_fixtures(directory: Path, rows: int, columns: int) -> list[ExcelFixture]:
    if not 2 <= rows <= 1_048_576 or not 2 <= columns <= 16_384:
        raise ValueError("rows/columns must be within XLSX limits and >= 2")
    directory.mkdir(parents=True, exist_ok=True)
    fixtures = []
    for kind in ("dense", "sparse", "formula", "dates"):
        path = directory / f"{kind}.xlsx"
        first_cell = "A1"
        last_cell = "ALL100000" if kind == "sparse" else f"{get_column_letter(columns)}{rows}"
        first_query, last_query, both_query = FIRST, LAST, "sf_excel_hit"
        if kind == "sparse":
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet[first_cell], sheet[last_cell] = FIRST, LAST
        else:
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("Data")
            for row in range(1, rows + 1):
                if kind == "formula":
                    values = ["=1+1"] * columns
                elif kind == "dates":
                    values = [datetime(2025, 1, 1, 12, 0)] * columns
                else:
                    values = [f"payload-{row}-{col}" for col in range(columns)]
                if row == 1:
                    values[0] = '="sf_excel_hit_first"' if kind == "formula" else FIRST
                if row == rows:
                    values[-1] = '="sf_excel_hit_last"' if kind == "formula" else LAST
                if kind == "dates":
                    first_query, last_query, both_query = "2026-10-04 01:02:03", "2026-10-04 04:05:06", "2026-10-04"
                    if row == 1:
                        values[0] = datetime(2026, 10, 4, 1, 2, 3)
                    if row == rows:
                        values[-1] = datetime(2026, 10, 4, 4, 5, 6)
                sheet.append(values)
        workbook.save(path)
        workbook.close()
        if kind == "formula":
            _cache_formulas(path, first_cell, last_cell)
        fixtures.append(ExcelFixture(
            kind, path, first_query, last_query, both_query, first_cell, last_cell,
            2 if kind == "sparse" else rows * columns,
            100_000_000 if kind == "sparse" else rows * columns,
        ))
    return fixtures


def cases_for(fixture: ExcelFixture) -> list[ExcelCase]:
    cases = []
    for precise in (False, True):
        for existence in (False, True):
            for name, query, cells in (
                ("first", fixture.first_query, [fixture.first_cell]),
                ("last", fixture.last_query, [fixture.last_cell]),
                ("both", fixture.both_query, [fixture.first_cell, fixture.last_cell]),
                ("absent", "sf_excel_absent_6b18f043", []),
            ):
                profile = "precise" if precise else "normal"
                mode = "existence" if existence else "full"
                cases.append(ExcelCase(f"Excel {fixture.name}/{profile}/{mode}/{name}",
                                       query, existence, frozenset(cells), precise))
    return cases


def _run(path: Path, case: ExcelCase):
    from core.search_engine import search_in_excel_special

    return search_in_excel_special(str(path), case.query, use_complex_search=case.precise,
                                   existence_only=case.existence_only)


def validate_result(result, case: ExcelCase) -> None:
    expected_count = min(1, len(case.expected_cells)) if case.existence_only else len(case.expected_cells)
    if expected_count == 0:
        if result is not None:
            raise RuntimeError(f"{case.name}: expected no result, got {result!r}")
        return
    if not result or len(result) != 3 or result[1] != expected_count:
        raise RuntimeError(f"{case.name}: unexpected result {result!r}")
    matches = result[2]
    if len(matches) != expected_count or any(match[0] < 0 for match in matches):
        raise RuntimeError(f"{case.name}: skipped/partial/duplicate results: {matches!r}")
    if not case.existence_only:
        coordinates = [(match[1], match[2]) for match in matches]
        expected = {("Data", cell) for cell in case.expected_cells}
        if set(coordinates) != expected or len(coordinates) != len(expected):
            raise RuntimeError(f"{case.name}: wrong coordinates {coordinates!r}")
        if any(case.query.casefold() not in str(match[3]).casefold() for match in matches):
            raise RuntimeError(f"{case.name}: wrong cell values {matches!r}")


def measure(path: Path, case: ExcelCase, repeats: int) -> dict:
    if repeats < 1:
        raise ValueError("repeats must be positive")
    validate_result(_run(path, case), case)  # Excluded warmup per case.
    samples = []
    rss_samples = []
    process = psutil.Process()
    peak = [process.memory_info().rss / 1024**2]
    stopped = threading.Event()

    def sample_memory():
        while not stopped.wait(0.01):
            peak[0] = max(peak[0], process.memory_info().rss / 1024**2)

    sampler = threading.Thread(target=sample_memory, daemon=True)
    sampler.start()
    try:
        for _ in range(repeats):
            gc.collect()
            peak[0] = process.memory_info().rss / 1024**2
            started = time.perf_counter()
            result = _run(path, case)
            elapsed = time.perf_counter() - started
            validate_result(result, case)
            samples.append(elapsed)
            peak[0] = max(peak[0], process.memory_info().rss / 1024**2)
            rss_samples.append(peak[0])
    finally:
        stopped.set()
        sampler.join()
    return {"dataset": case.name, "samples": samples,
            "median_seconds": statistics.median(samples),
            "min_seconds": min(samples), "max_seconds": max(samples),
            "rss_samples_mib": rss_samples, "peak_rss_mib": max(rss_samples),
            "hits": min(1, len(case.expected_cells)) if case.existence_only else len(case.expected_cells)}


def run_suite(directory: Path, rows: int, columns: int, repeats: int) -> dict:
    from core.search_engine import HAS_RUST_ENGINE
    from rust_engine import sf_engine
    from sf_utils._version import VERSION

    if not HAS_RUST_ENGINE or sf_engine.ENGINE_VERSION != VERSION:
        raise RuntimeError("A matching project Rust engine is required; rebuild before measuring")
    fixtures = create_fixtures(directory, rows, columns)
    results = []
    for fixture in fixtures:
        for case in cases_for(fixture):
            result = measure(fixture.path, case, repeats)
            results.append(result)
            print(f"{case.name}: {result['median_seconds']:.4f}s "
                  f"[{result['min_seconds']:.4f}, {result['max_seconds']:.4f}], "
                  f"hits={result['hits']}", flush=True)
    engine = ROOT / "src/rust_engine/sf_engine.pyd"
    return {"version": VERSION, "engine_sha256": hashlib.sha256(engine.read_bytes()).hexdigest(),
            "environment": {"platform": platform.platform(), "python": platform.python_version(),
                            "logical_cpus": psutil.cpu_count(),
                            "ram_bytes": psutil.virtual_memory().total},
            "rows": rows, "columns": columns, "repeats": repeats,
            "scope": "in-process Excel function; excludes GUI, worker startup and formula calculation",
            "fixtures": [{"name": f.name, "stored_cells": f.stored_cells,
                          "rectangle_cells": f.rectangle_cells, "size_bytes": f.path.stat().st_size,
                          "sha256": hashlib.sha256(f.path.read_bytes()).hexdigest()} for f in fixtures],
            "results": results}


def record_history(report: dict, tag: str) -> None:
    from scripts.benchmark_performance import save_benchmark_history

    for index in range(report["repeats"]):
        # Synchronous API: latency is completion time, not first streamed result.
        save_benchmark_history([
            {"dataset": r["dataset"], "total_time": r["samples"][index],
             "latency": r["samples"][index], "jitter": 0.0,
             "peak_rss": r["rss_samples_mib"][index], "results_count": r["hits"], "skipped_count": 0}
            for r in report["results"]
        ], tag=f"{tag}-r{index + 1}")


def close_owned_logs(directory: Path) -> None:
    """Release only this CLI's temporary logs before Windows directory cleanup."""
    loggers = [logging.getLogger(), *logging.Logger.manager.loggerDict.values()]
    for logger in loggers:
        if not isinstance(logger, logging.Logger):
            continue
        for handler in list(logger.handlers):
            if isinstance(handler, logging.FileHandler) and Path(handler.baseFilename).resolve().is_relative_to(directory.resolve()):
                logger.removeHandler(handler)
                handler.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, default=25_000)
    parser.add_argument("--columns", type=int, default=4)
    parser.add_argument("--repeats", type=int, default=5)
    parser.add_argument("--output", type=Path, help="Optional JSON measurement report")
    parser.add_argument("--record-tag", help="Append validated measurements to benchmark_history.md")
    args = parser.parse_args()
    if not 2 <= args.rows <= 1_048_576 or not 2 <= args.columns <= 16_384 or args.repeats < 1:
        parser.error("XLSX rows/columns must be within limits and >= 2; repeats must be positive")
    # Import search code only after isolating defaults from the user's settings.
    with tempfile.TemporaryDirectory(prefix="stringfinder-excel-suite-") as temp_dir:
        previous_appdata = os.environ.get("APPDATA")
        try:
            os.environ["APPDATA"] = str(Path(temp_dir) / "appdata")
            report = run_suite(Path(temp_dir) / "fixtures", args.rows, args.columns, args.repeats)
            if args.output:
                args.output.parent.mkdir(parents=True, exist_ok=True)
                args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            if args.record_tag:
                record_history(report, args.record_tag)
        finally:
            close_owned_logs(Path(temp_dir))
            if previous_appdata is None:
                os.environ.pop("APPDATA", None)
            else:
                os.environ["APPDATA"] = previous_appdata


if __name__ == "__main__":
    main()
